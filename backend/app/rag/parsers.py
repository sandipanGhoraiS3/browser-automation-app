from __future__ import annotations

import csv
import io
import json
import re
from pathlib import Path
from typing import Protocol
from app.rag.models import Block, ParsedDocument


class UnsupportedDocument(ValueError):
    pass


class EmptyDocument(ValueError):
    pass


class OCRRequired(ValueError):
    pass


class DocumentParser(Protocol):
    def parse(self, data: bytes, file_name: str) -> ParsedDocument: ...


def _clean(text: str) -> str:
    text = text.replace('\x00', '').replace('\r\n', '\n').replace('\r', '\n')
    return re.sub(r'[ \t]+', ' ', text).strip()


def _looks_heading(text: str) -> bool:
    line = text.strip().lstrip('#').strip()
    return bool(line and len(line) <= 120 and (
        text.lstrip().startswith('#')
        or re.match(r'^\d+(?:\.\d+)*[.)]?\s+\S+', line)
        or (line.upper() == line and any(c.isalpha() for c in line))
    ))


def _text_blocks(text: str, markdown: bool = False) -> list[Block]:
    blocks: list[Block] = []
    section = None
    paragraphs = re.split(r'\n\s*\n+', _clean(text))
    for paragraph in paragraphs:
        paragraph = paragraph.strip()
        if not paragraph:
            continue
        first = paragraph.splitlines()[0].strip()
        if _looks_heading(first) or (markdown and first.startswith('#')):
            section = first.lstrip('#').strip()
            if len(paragraph.splitlines()) == 1:
                continue
        blocks.append(Block(paragraph, section=section, kind='paragraph'))
    return blocks


class TextParser:
    def parse(self, data: bytes, file_name: str) -> ParsedDocument:
        try:
            text = data.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = data.decode('utf-8', errors='replace')
        blocks = _text_blocks(text, Path(file_name).suffix.lower() == '.md')
        if not blocks:
            raise EmptyDocument('The document contains no readable text.')
        return ParsedDocument(blocks)


class PdfParser:
    def parse(self, data: bytes, file_name: str) -> ParsedDocument:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        blocks: list[Block] = []
        for number, page in enumerate(reader.pages, 1):
            text = _clean(page.extract_text() or '')
            if not text:
                continue
            section = None
            for paragraph in re.split(r'\n\s*\n+|(?<=\.)\n(?=[A-Z])', text):
                paragraph = paragraph.strip()
                if not paragraph:
                    continue
                first = paragraph.splitlines()[0]
                if _looks_heading(first):
                    section = first.strip()
                blocks.append(Block(paragraph, page_number=number, section=section))
        if not blocks:
            raise OCRRequired('The PDF has no extractable text. OCR is required.')
        return ParsedDocument(blocks, page_count=len(reader.pages))


class DocxParser:
    def parse(self, data: bytes, file_name: str) -> ParsedDocument:
        from docx import Document
        document = Document(io.BytesIO(data))
        blocks: list[Block] = []
        section = None
        for paragraph in document.paragraphs:
            text = _clean(paragraph.text)
            if not text:
                continue
            if paragraph.style and paragraph.style.name.lower().startswith('heading'):
                section = text
            else:
                blocks.append(Block(text, section=section))
        for table in document.tables:
            rows = [' | '.join(_clean(cell.text) for cell in row.cells) for row in table.rows]
            table_text = '\n'.join(row for row in rows if row.strip(' |'))
            if table_text:
                blocks.append(Block(table_text, section=section, kind='table'))
        if not blocks:
            raise EmptyDocument('The document contains no readable text.')
        return ParsedDocument(blocks)


class CsvParser:
    def parse(self, data: bytes, file_name: str) -> ParsedDocument:
        text = data.decode('utf-8-sig', errors='replace')
        rows = list(csv.reader(io.StringIO(text)))
        if not rows:
            raise EmptyDocument('The CSV contains no rows.')
        header = rows[0]
        blocks = []
        for index, row in enumerate(rows[1:] or rows, 1):
            if rows[1:]:
                rendered = ' | '.join(f'{header[i] if i < len(header) else f"column_{i + 1}"}: {value}' for i, value in enumerate(row))
            else:
                rendered = ' | '.join(row)
            if rendered.strip(' |'):
                blocks.append(Block(rendered, section=f'Rows {index}', kind='table'))
        if not blocks:
            raise EmptyDocument('The CSV contains no readable values.')
        return ParsedDocument(blocks)


class XlsxParser:
    def parse(self, data: bytes, file_name: str) -> ParsedDocument:
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        blocks: list[Block] = []
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue
            header = [str(value) if value is not None else f'column_{i + 1}' for i, value in enumerate(rows[0])]
            for index, row in enumerate(rows[1:] or rows, 1):
                rendered = ' | '.join(f'{header[i] if i < len(header) else f"column_{i + 1}"}: {value}' for i, value in enumerate(row) if value is not None)
                if rendered:
                    blocks.append(Block(rendered, section=sheet.title, kind='table'))
        if not blocks:
            raise EmptyDocument('The workbook contains no readable values.')
        return ParsedDocument(blocks)


class JsonParser:
    def parse(self, data: bytes, file_name: str) -> ParsedDocument:
        try:
            value = json.loads(data.decode('utf-8-sig'))
        except (UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise ValueError('The JSON document is invalid.') from exc
        blocks: list[Block] = []
        entries = value if isinstance(value, list) else [value]
        for index, entry in enumerate(entries, 1):
            text = json.dumps(entry, ensure_ascii=False, indent=2)
            blocks.append(Block(text, section=f'Item {index}', kind='structured'))
        if not blocks:
            raise EmptyDocument('The JSON document is empty.')
        return ParsedDocument(blocks)


PARSERS: dict[str, DocumentParser] = {
    '.pdf': PdfParser(), '.docx': DocxParser(), '.txt': TextParser(), '.md': TextParser(),
    '.csv': CsvParser(), '.xlsx': XlsxParser(), '.json': JsonParser(),
}


def parse_document(data: bytes, file_name: str) -> ParsedDocument:
    suffix = Path(file_name).suffix.lower()
    parser = PARSERS.get(suffix)
    if not parser:
        raise UnsupportedDocument('Supported formats: PDF, DOCX, TXT, MD, CSV, XLSX, JSON.')
    return parser.parse(data, file_name)
