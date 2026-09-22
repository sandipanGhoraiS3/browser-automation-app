import io
import asyncio
import json
import math
from types import SimpleNamespace
import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from app.database.store import DEFAULT_WORKSPACE_ID, store as app_store
from app.rag.chunking import StructureAwareChunker
from app.rag.parsers import OCRRequired, parse_document
from app.rag.service import RagService


class FakeBlobStore:
    def __init__(self):
        self.values = {}

    async def upload(self, path, data, content_type):
        if path in self.values:
            raise RuntimeError('blob already exists')
        self.values[path] = (data, content_type)

    async def download(self, path):
        return self.values[path][0]

    async def delete(self, path):
        self.values.pop(path, None)

    async def health(self):
        return None

    async def initialize(self):
        return {'container': 'test', 'created': False}


class FakeEmbeddings:
    available = True

    @staticmethod
    def vector(text):
        lowered = text.lower()
        values = [lowered.count(term) for term in ('finance', 'reimbursement', 'hr', 'leave')]
        norm = math.sqrt(sum(value * value for value in values)) or 1
        return [value / norm for value in values]

    async def embed(self, texts):
        return [self.vector(text) for text in texts]

    async def health(self):
        return None


class FakeVectorStore:
    def __init__(self):
        self.values = {}

    async def health(self):
        return None

    async def initialize(self):
        return {'partition_key': '/workspace_id', 'vector_path': '/embedding'}

    async def upsert(self, item):
        self.values[(item['workspace_id'], item['id'])] = dict(item)

    async def get(self, item_id, workspace_id):
        value = self.values.get((workspace_id, item_id))
        return dict(value) if value else None

    async def delete(self, item_id, workspace_id):
        self.values.pop((workspace_id, item_id), None)

    async def query_items(self, query, parameters, workspace_id):
        params = {item['name']: item['value'] for item in parameters}
        rows = [dict(item) for (partition, _), item in self.values.items() if partition == workspace_id]
        for kind in ('domain', 'document', 'chunk'):
            if f"kind='{kind}'" in query:
                rows = [item for item in rows if item.get('kind') == kind]
        for name in ('domain_id', 'document_id', 'sha256'):
            value = params.get('@' + name)
            if value:
                rows = [item for item in rows if item.get(name) == value]
        if '@name' in params:
            rows = [item for item in rows if item.get('name', '').lower() == params['@name']]
        return rows[:1] if 'TOP 1' in query else rows

    async def vector_search(self, workspace_id, embedding, limit, filters):
        rows = []
        for (partition, _), item in self.values.items():
            if partition != workspace_id or item.get('kind') != 'chunk':
                continue
            if any(value and item.get(name) != value for name, value in filters.items()):
                continue
            dot = sum(a * b for a, b in zip(embedding, item['embedding']))
            rows.append({**item, 'distance': 1 - dot})
        return sorted(rows, key=lambda row: row['distance'])[:limit]

    async def delete_document_items(self, workspace_id, document_id):
        for key, item in list(self.values.items()):
            if key[0] == workspace_id and item.get('document_id') == document_id:
                del self.values[key]


def test_supported_parsers_and_structure_aware_chunking():
    markdown = parse_document(b'# Policy\n\nFinance reimbursement applies.\n\n## Limits\n\nReceipts are required.', 'policy.md')
    assert markdown.blocks[0].section == 'Policy'
    assert markdown.blocks[-1].section == 'Limits'
    assert StructureAwareChunker(50, 10).chunk(markdown)[0].chunk_hash

    assert parse_document(b'plain enterprise text', 'notes.txt').blocks
    assert parse_document(b'name,policy\nA,reimbursement', 'rows.csv').blocks
    assert parse_document(json.dumps({'finance': 'reimbursement'}).encode(), 'data.json').blocks

    docx = Document()
    docx.add_heading('Finance', 1)
    docx.add_paragraph('Reimbursement requires a receipt.')
    docx_bytes = io.BytesIO()
    docx.save(docx_bytes)
    assert parse_document(docx_bytes.getvalue(), 'policy.docx').blocks[0].section == 'Finance'

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Finance'
    sheet.append(['Policy', 'Rule'])
    sheet.append(['Reimbursement', 'Receipt required'])
    xlsx_bytes = io.BytesIO()
    workbook.save(xlsx_bytes)
    assert parse_document(xlsx_bytes.getvalue(), 'policy.xlsx').blocks[0].section == 'Finance'

    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    pdf_bytes = io.BytesIO()
    writer.write(pdf_bytes)
    with pytest.raises(OCRRequired):
        parse_document(pdf_bytes.getvalue(), 'scan.pdf')


@pytest.mark.asyncio
async def test_ingestion_duplicate_reindex_delete_and_health():
    vectors, blobs = FakeVectorStore(), FakeBlobStore()
    service = RagService(vectors, blobs, FakeEmbeddings())
    domain = await service.create_domain(DEFAULT_WORKSPACE_ID, 'Finance')
    document = await service.upload(
        DEFAULT_WORKSPACE_ID, domain['id'], 'policy.txt', 'text/plain',
        b'Finance reimbursement policy requires a receipt. ' * 20,
    )
    await service.tasks.copy().pop()
    ready = await service.get_document(DEFAULT_WORKSPACE_ID, document['id'])
    assert ready['status'] == 'READY'
    assert ready['chunk_count'] >= 1
    assert ready['blob_path'] in blobs.values
    with pytest.raises(ValueError, match='Duplicate'):
        await service.upload(
            DEFAULT_WORKSPACE_ID, domain['id'], 'copy.txt', 'text/plain',
            b'Finance reimbursement policy requires a receipt. ' * 20,
        )
    await service.reindex(DEFAULT_WORKSPACE_ID, document['id'])
    await service.tasks.copy().pop()
    assert (await service.get_document(DEFAULT_WORKSPACE_ID, document['id']))['status'] == 'READY'
    health = await service.health()
    assert all(item['status'] == 'connected' for item in health['services'].values())
    await service.delete_document(DEFAULT_WORKSPACE_ID, document['id'])
    assert await vectors.get(document['id'], DEFAULT_WORKSPACE_ID) is None
    assert not blobs.values


@pytest.mark.asyncio
async def test_interrupted_ingestion_is_recovered_from_blob():
    vectors, blobs = FakeVectorStore(), FakeBlobStore()
    service = RagService(vectors, blobs, FakeEmbeddings())
    domain = await service.create_domain(DEFAULT_WORKSPACE_ID, 'Recovery')
    document = await service.upload(
        DEFAULT_WORKSPACE_ID, domain['id'], 'recovery.txt', 'text/plain',
        b'Finance reimbursement recovery content. ' * 20,
    )
    task = next(iter(service.tasks))
    task.cancel()
    await asyncio.gather(task, return_exceptions=True)
    interrupted = await service.get_document(DEFAULT_WORKSPACE_ID, document['id'])
    interrupted.update(status='INDEXING', progress=80)
    await vectors.upsert(interrupted)

    assert await service.recover_interrupted() == 1
    await next(iter(service.tasks))
    recovered = await service.get_document(DEFAULT_WORKSPACE_ID, document['id'])
    assert recovered['status'] == 'READY'
    assert recovered['progress'] == 100


@pytest.mark.asyncio
async def test_delete_domain_removes_documents_chunks_blob_and_domain():
    vectors, blobs = FakeVectorStore(), FakeBlobStore()
    service = RagService(vectors, blobs, FakeEmbeddings())
    domain = await service.create_domain(DEFAULT_WORKSPACE_ID, 'Delete Me')
    document = await service.upload(
        DEFAULT_WORKSPACE_ID, domain['id'], 'delete.txt', 'text/plain',
        b'Finance reimbursement deletion content. ' * 20,
    )
    await next(iter(service.tasks))

    result = await service.delete_domain(DEFAULT_WORKSPACE_ID, domain['id'])

    assert result == {'ok': True, 'deleted_documents': 1}
    assert await vectors.get(domain['id'], DEFAULT_WORKSPACE_ID) is None
    assert await vectors.get(document['id'], DEFAULT_WORKSPACE_ID) is None
    assert not blobs.values


@pytest.mark.asyncio
async def test_retrieval_enforces_workspace_and_explicit_domain_isolation():
    vectors = FakeVectorStore()
    service = RagService(vectors, FakeBlobStore(), FakeEmbeddings())
    workspace_b = app_store.create_workspace('RAG Isolation Test')['id']
    finance = await service.create_domain(DEFAULT_WORKSPACE_ID, 'Finance Isolation')
    hr = await service.create_domain(DEFAULT_WORKSPACE_ID, 'HR Isolation')
    other = await service.create_domain(workspace_b, 'Other Workspace')

    async def chunk(workspace_id, domain_id, item_id, text):
        await vectors.upsert({
            'id': item_id, 'kind': 'chunk', 'workspace_id': workspace_id,
            'domain_id': domain_id, 'document_id': 'doc-' + item_id,
            'chunk_index': 0, 'text': text, 'embedding': FakeEmbeddings.vector(text),
            'page_number': 2, 'section': 'Policy', 'source_file': item_id + '.txt',
            'content_type': 'text/plain',
        })

    await chunk(DEFAULT_WORKSPACE_ID, finance['id'], 'finance', 'finance reimbursement receipt policy')
    await chunk(DEFAULT_WORKSPACE_ID, hr['id'], 'hr', 'hr leave policy')
    await chunk(workspace_b, other['id'], 'leak', 'finance reimbursement secret from workspace b')

    result = await service.query(DEFAULT_WORKSPACE_ID, 'finance reimbursement', finance['id'])
    assert [item['id'] for item in result.chunks] == ['finance']
    assert result.sources[0]['page_number'] == 2
    assert 'leak' not in {item['id'] for item in result.chunks}

    no_cross_domain = await service.query(DEFAULT_WORKSPACE_ID, 'hr leave', finance['id'])
    assert 'hr' not in {item['id'] for item in no_cross_domain.chunks}
